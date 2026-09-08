"""
Multi-tool Ollama agent
=======================
One tool-calling loop wired up to three capability groups:
  1. Web tools    - search + fetch/scrape pages (Ollama's own hosted API)
  2. File tools   - write plain-text formats (.md/.txt/.json/.csv/.html/.py)
                     plus Word (.docx) and Excel (.xlsx) if those libs are installed
  3. CLI tool     - run whitelisted shell commands, with human confirmation

Setup:
    pip install ollama python-docx openpyxl

    Set your Ollama API key (only needed for web_search / web_fetch):
      macOS/Linux (bash/zsh):  export OLLAMA_API_KEY=your_key_here
      Windows PowerShell:      $env:OLLAMA_API_KEY = "your_key_here"
      Windows (persist it):    setx OLLAMA_API_KEY "your_key_here"   (new terminal needed after)

Run:
    python ollama_agent.py
"""

import os
import platform
import re
import shlex
import shutil
import subprocess
from pathlib import Path

from ollama import chat, web_fetch, web_search

# --------------------------------------------------------------------------
# Safety configuration for the CLI tool. Edit this list to fit your workflow
# -- keep it as short as you can. Nothing outside this list will ever run.
#
# NOTE: ls/dir/cat/pwd/echo/grep/mkdir/cp/mv are NOT real executables on
# Windows (they're PowerShell aliases or cmd.exe built-ins), so subprocess
# can never find an "ls.exe". They're implemented natively in Python below
# instead (see NATIVE_COMMANDS) so they work the same on every OS.
# ALLOWED_COMMANDS is only for real external programs on PATH.
# --------------------------------------------------------------------------
IS_WINDOWS = platform.system() == "Windows"

ALLOWED_COMMANDS = {"python", "python3", "pip", "git", "curl"}
CONFIRM_BEFORE_RUN = True   # ask a human before executing anything
COMMAND_TIMEOUT = 30        # seconds
MAX_OUTPUT_CHARS = 4000     # keep tool output from flooding the model's context


# --------------------------------------------------------------------------
# Native, shell-free implementations of common filesystem commands.
# Each takes (args, stdin_text) so they can sit anywhere in a pipeline --
# stdin_text is the previous stage's output when chained with '|'.
# --------------------------------------------------------------------------
def _native_ls(args, stdin_text=None):
    # This is a simplified listing, not a full ls/dir reimplementation --
    # flags like -l, -la, -a are ignored rather than causing an error.
    # The first non-flag token (if any) is used as the target path.
    positional = [a for a in args if not a.startswith("-")]
    target = positional[0] if positional else "."
    p = Path(target)
    if not p.exists():
        return f"No such path: {target}"
    if p.is_file():
        return p.name
    return "\n".join(sorted(x.name + ("/" if x.is_dir() else "") for x in p.iterdir()))

def _native_pwd(args, stdin_text=None):
    return os.getcwd()

def _native_cat(args, stdin_text=None):
    if not args:
        return "cat: missing filename"
    try:
        return Path(args[0]).read_text(encoding="utf-8", errors="replace")[:MAX_OUTPUT_CHARS]
    except Exception as e:
        return f"cat: {e}"

def _native_echo(args, stdin_text=None):
    return " ".join(args)

def _native_grep(args, stdin_text=None):
    if not args:
        return "grep: missing pattern"
    pattern = args[0]
    if stdin_text is not None:
        lines = stdin_text.splitlines()
    elif len(args) > 1:
        try:
            lines = Path(args[1]).read_text(encoding="utf-8", errors="replace").splitlines()
        except Exception as e:
            return f"grep: {e}"
    else:
        return "grep: no input (pipe something into it, or pass a filename)"
    try:
        rx = re.compile(pattern)
    except re.error as e:
        return f"grep: bad pattern: {e}"
    matches = [line for line in lines if rx.search(line)]
    return "\n".join(matches) if matches else "(no matches)"

def _native_mkdir(args, stdin_text=None):
    if not args:
        return "mkdir: missing path"
    Path(args[0]).mkdir(parents=True, exist_ok=True)
    return f"Created {Path(args[0]).resolve()}"

def _native_cp(args, stdin_text=None):
    if len(args) < 2:
        return "cp: usage: cp <source> <dest>"
    shutil.copy2(args[0], args[1])
    return f"Copied {args[0]} -> {args[1]}"

def _native_mv(args, stdin_text=None):
    if len(args) < 2:
        return "mv: usage: mv <source> <dest>"
    shutil.move(args[0], args[1])
    return f"Moved {args[0]} -> {args[1]}"

NATIVE_COMMANDS = {
    "ls": _native_ls, "dir": _native_ls,
    "pwd": _native_pwd,
    "cat": _native_cat, "type": _native_cat,
    "echo": _native_echo,
    "grep": _native_grep,
    "mkdir": _native_mkdir,
    "cp": _native_cp,
    "mv": _native_mv,
}


# --------------------------------------------------------------------------
# Tool 1: CLI execution (guarded)
# --------------------------------------------------------------------------
def _run_segment(tokens: list, stdin_text) -> str:
    program = tokens[0]
    if program in NATIVE_COMMANDS:
        return NATIVE_COMMANDS[program](tokens[1:], stdin_text)
    result = subprocess.run(
        tokens, shell=False, capture_output=True, text=True,
        input=stdin_text, timeout=COMMAND_TIMEOUT, cwd=os.getcwd(),
    )
    return result.stdout + result.stderr


def run_command(command: str) -> str:
    """Run a command, or a pipeline of commands separated by '|'.

    Example: "git log --oneline | grep fix"

    No shell is ever invoked. The string is split only on the literal '|'
    character, each stage is tokenized with shlex, and each stage's program
    name must be in ALLOWED_COMMANDS or NATIVE_COMMANDS or the whole thing
    is blocked before anything runs. Because there's no real shell, other
    shell syntax (&&, ;, backticks, $(), >, >>) is not treated as an
    operator -- it's just inert text passed as a literal argument.

    Residual risk: a program you DO allow can still do dangerous things
    with its own flags (e.g. "python -c ...", "curl | some-installer").
    The allowlist only stops programs outside it -- keep it short and
    review it, especially before adding anything that can itself execute
    arbitrary code.
    """
    segments = [seg.strip() for seg in command.split("|") if seg.strip()]
    if not segments:
        return "Empty command."

    parsed = []
    for seg in segments:
        try:
            parsed.append(shlex.split(seg, posix=not IS_WINDOWS))
        except ValueError as e:
            return f"Could not parse '{seg}': {e}"

    known = ALLOWED_COMMANDS | set(NATIVE_COMMANDS)
    blocked = sorted({t[0] for t in parsed if t and t[0] not in known})
    if blocked:
        return f"Blocked: {', '.join(blocked)} not allowed. Allowed: {', '.join(sorted(known))}"

    if CONFIRM_BEFORE_RUN:
        approved = input(f"\n[agent wants to run] {command}\nAllow? (y/n): ").strip().lower()
        if approved != "y":
            return "Command was not approved by the user."

    try:
        stdin_text, output = None, ""
        for tokens in parsed:
            output = _run_segment(tokens, stdin_text)
            stdin_text = output
        output = output.strip()
        return output[:MAX_OUTPUT_CHARS] if output else "(command produced no output)"
    except subprocess.TimeoutExpired:
        return f"Command timed out after {COMMAND_TIMEOUT}s."
    except FileNotFoundError as e:
        return f"Program not found: {e}"
    except Exception as e:
        return f"Error running command: {e}"


# --------------------------------------------------------------------------
# Tool 2: file creation
# --------------------------------------------------------------------------
def write_file(path: str, content: str) -> str:
    """Create or overwrite a plain text-based file.

    Works for any text format: .txt, .md, .json, .csv, .html, .py, .yaml, etc.
    `content` is written exactly as given, so ask the model to produce
    correctly-formatted content (valid JSON, valid CSV rows, etc.) up front.
    """
    p = Path(path).expanduser()
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(content, encoding="utf-8")
    return f"Wrote {len(content)} characters to {p.resolve()}"


def write_docx(path: str, content: str) -> str:
    """Create a Word document. `content` is plain text; each line becomes a paragraph."""
    try:
        from docx import Document
    except ImportError:
        return "python-docx is not installed. Run: pip install python-docx"

    p = Path(path).expanduser()
    p.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for line in content.split("\n"):
        doc.add_paragraph(line)
    doc.save(p)
    return f"Wrote Word document to {p.resolve()}"


def write_xlsx(path: str, rows_csv: str) -> str:
    """Create an Excel file from CSV-style text (one row per line, comma-separated cells)."""
    try:
        from openpyxl import Workbook
    except ImportError:
        return "openpyxl is not installed. Run: pip install openpyxl"

    p = Path(path).expanduser()
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    for line in rows_csv.strip().split("\n"):
        ws.append(line.split(","))
    wb.save(p)
    return f"Wrote spreadsheet to {p.resolve()}"


# --------------------------------------------------------------------------
# Tool registry + main loop
# --------------------------------------------------------------------------
available_tools = {
    "web_search": web_search,
    "web_fetch": web_fetch,
    "write_file": write_file,
    "write_docx": write_docx,
    "write_xlsx": write_xlsx,
    "run_command": run_command,
}
tools = list(available_tools.values())


def agent(user_message: str, model: str = "qwen3:8b") -> None:
    messages = [{"role": "user", "content": user_message}]

    while True:
        response = chat(model=model, messages=messages, tools=tools, think=True)

        if response.message.content:
            print("\nAgent:", response.message.content)
        messages.append(response.message)

        if not response.message.tool_calls:
            break  # model gave a final answer, stop looping

        for call in response.message.tool_calls:
            fn = available_tools.get(call.function.name)
            if fn is None:
                result = f"Tool '{call.function.name}' not found."
            else:
                try:
                    result = fn(**call.function.arguments)
                except Exception as e:
                    result = f"Tool error: {e}"

            print(f"[tool] {call.function.name}({call.function.arguments}) -> {str(result)[:200]}")
            messages.append({
                "role": "tool",
                "content": str(result)[:MAX_OUTPUT_CHARS],
                "tool_name": call.function.name,
            })


if __name__ == "__main__":
    agent(
        "Search for the latest Ollama release notes, save a short summary "
        "as ollama_notes.md, then list the files in the current directory."
    )

    agent(
            "Generate a very-very unique and technical knowledge base on "
            "making the AI work like GOD, and save the content in very "
            "beautiful and elegane formatted structure "
            "in the temp/AI_Advanced_USE.md, then list the files in the temp directory."
        )


"""
More requirements:
python-docx
openpyxl
"""

