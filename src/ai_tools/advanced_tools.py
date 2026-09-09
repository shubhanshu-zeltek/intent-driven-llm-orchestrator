import os
import platform
import re
import shlex
import shutil
import subprocess
from pathlib import Path
from langchain_core.tools import tool, StructuredTool
from ollama import web_fetch, web_search


IS_WINDOWS = platform.system() == "Windows"

ALLOWED_COMMANDS = {"python", "python3", "pip", "git", "curl"}
COMMAND_TIMEOUT = 30        # seconds
MAX_OUTPUT_CHARS = 4000     # keep tool output from flooding the model's context


# --------------------------------------------------------------------------
# Native, shell-free implementations of common filesystem commands.
# Each takes (args, stdin_text) so they can sit anywhere in a pipeline --
# stdin_text is the previous stage's output when chained with '|'.
# --------------------------------------------------------------------------
def _native_ls(args, stdin_text=None):
    # This is a simplified listing, not a full ls/dir reimplementation --
    # flags like -l, -la, -a are ignored rather than causing an error.
    # The first non-flag token (if any) is used as the target path.
    positional = [a for a in args if not a.startswith("-")]
    target = positional[0] if positional else "."
    p = Path(target)
    if not p.exists():
        return f"No such path: {target}"
    if p.is_file():
        return p.name
    return "\n".join(sorted(x.name + ("/" if x.is_dir() else "") for x in p.iterdir()))

def _native_pwd(args, stdin_text=None):
    return os.getcwd()

def _native_cat(args, stdin_text=None):
    if not args:
        return "cat: missing filename"
    try:
        return Path(args[0]).read_text(encoding="utf-8", errors="replace")[:MAX_OUTPUT_CHARS]
    except Exception as e:
        return f"cat: {e}"

def _native_echo(args, stdin_text=None):
    return " ".join(args)

def _native_grep(args, stdin_text=None):
    if not args:
        return "grep: missing pattern"
    pattern = args[0]
    if stdin_text is not None:
        lines = stdin_text.splitlines()
    elif len(args) > 1:
        try:
            lines = Path(args[1]).read_text(encoding="utf-8", errors="replace").splitlines()
        except Exception as e:
            return f"grep: {e}"
    else:
        return "grep: no input (pipe something into it, or pass a filename)"
    try:
        rx = re.compile(pattern)
    except re.error as e:
        return f"grep: bad pattern: {e}"
    matches = [line for line in lines if rx.search(line)]
    return "\n".join(matches) if matches else "(no matches)"

def _native_mkdir(args, stdin_text=None):
    if not args:
        return "mkdir: missing path"
    Path(args[0]).mkdir(parents=True, exist_ok=True)
    return f"Created {Path(args[0]).resolve()}"

def _native_cp(args, stdin_text=None):
    if len(args) < 2:
        return "cp: usage: cp <source> <dest>"
    shutil.copy2(args[0], args[1])
    return f"Copied {args[0]} -> {args[1]}"

def _native_mv(args, stdin_text=None):
    if len(args) < 2:
        return "mv: usage: mv <source> <dest>"
    shutil.move(args[0], args[1])
    return f"Moved {args[0]} -> {args[1]}"

NATIVE_COMMANDS = {
    "ls": _native_ls, "dir": _native_ls,
    "pwd": _native_pwd,
    "cat": _native_cat, "type": _native_cat,
    "echo": _native_echo,
    "grep": _native_grep,
    "mkdir": _native_mkdir,
    "cp": _native_cp,
    "mv": _native_mv,
}


# --------------------------------------------------------------------------
# Tool 1: CLI execution (guarded)
# --------------------------------------------------------------------------
def _run_segment(tokens: list, stdin_text) -> str:
    program = tokens[0]
    if program in NATIVE_COMMANDS:
        return NATIVE_COMMANDS[program](tokens[1:], stdin_text)
    result = subprocess.run(
        tokens, shell=False, capture_output=True, text=True,
        input=stdin_text, timeout=COMMAND_TIMEOUT, cwd=os.getcwd(),
    )
    return result.stdout + result.stderr

@tool
def run_command(command: str) -> str:
    """Run a command, or a pipeline of commands separated by '|'.

    Example: "git log --oneline | grep fix"

    No shell is ever invoked. The string is split only on the literal '|'
    character, each stage is tokenized with shlex, and each stage's program
    name must be in ALLOWED_COMMANDS or NATIVE_COMMANDS or the whole thing
    is blocked before anything runs. Because there's no real shell, other
    shell syntax (&&, ;, backticks, $(), >, >>) is not treated as an
    operator -- it's just inert text passed as a literal argument.

    Residual risk: a program you DO allow can still do dangerous things
    with its own flags (e.g. "python -c ...", "curl | some-installer").
    The allowlist only stops programs outside it -- keep it short and
    review it, especially before adding anything that can itself execute
    arbitrary code.
    """
    segments = [seg.strip() for seg in command.split("|") if seg.strip()]
    if not segments:
        return "Empty command."

    parsed = []
    for seg in segments:
        try:
            parsed.append(shlex.split(seg, posix=not IS_WINDOWS))
        except ValueError as e:
            return f"Could not parse '{seg}': {e}"

    known = ALLOWED_COMMANDS | set(NATIVE_COMMANDS)
    blocked = sorted({t[0] for t in parsed if t and t[0] not in known})
    if blocked:
        return f"Blocked: {', '.join(blocked)} not allowed. Allowed: {', '.join(sorted(known))}"

    try:
        stdin_text, output = None, ""
        for tokens in parsed:
            output = _run_segment(tokens, stdin_text)
            stdin_text = output
        output = output.strip()
        return output[:MAX_OUTPUT_CHARS] if output else "(command produced no output)"
    except subprocess.TimeoutExpired:
        return f"Command timed out after {COMMAND_TIMEOUT}s."
    except FileNotFoundError as e:
        return f"Program not found: {e}"
    except Exception as e:
        return f"Error running command: {e}"


# --------------------------------------------------------------------------
# Tool 2: file creation
# --------------------------------------------------------------------------
@tool
def write_file(path: str, content: str) -> str:
    """Create or overwrite a plain text-based file.

    Works for any text format: .txt, .md, .json, .csv, .html, .py, .yaml, etc.
    `content` is written exactly as given, so ask the model to produce
    correctly-formatted content (valid JSON, valid CSV rows, etc.) up front.
    """
    p = Path(path).expanduser()
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(content, encoding="utf-8")
    return f"Wrote {len(content)} characters to {p.resolve()}"

@tool
def write_docx(path: str, content: str) -> str:
    """Create a Word document. `content` is plain text; each line becomes a paragraph."""
    try:
        from docx import Document
    except ImportError:
        return "python-docx is not installed. Run: pip install python-docx"

    p = Path(path).expanduser()
    p.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for line in content.split("\n"):
        doc.add_paragraph(line)
    doc.save(p)
    return f"Wrote Word document to {p.resolve()}"

@tool
def write_xlsx(path: str, rows_csv: str) -> str:
    """Create an Excel file from CSV-style text (one row per line, comma-separated cells)."""
    try:
        from openpyxl import Workbook
    except ImportError:
        return "openpyxl is not installed. Run: pip install openpyxl"

    p = Path(path).expanduser()
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    for line in rows_csv.strip().split("\n"):
        ws.append(line.split(","))
    wb.save(p)
    return f"Wrote spreadsheet to {p.resolve()}"


web_search_tool = StructuredTool.from_function(
    func=web_search,
    name="web_search",
    description=(web_search.__doc__ or "Search the web via Ollama's hosted search API.").strip(),
)
web_fetch_tool = StructuredTool.from_function(
    func=web_fetch,
    name="web_fetch",
    description=(web_fetch.__doc__ or "Fetch the content of a specific URL via Ollama's hosted fetch API.").strip(),
)

available_tools = {
    "web_search": web_search_tool,
    "web_fetch": web_fetch_tool,
    "write_file": write_file,
    "write_docx": write_docx,
    "write_xlsx": write_xlsx,
    "run_command": run_command,
}

def get_advanced_tools_name():
    return list(available_tools)

def get_advanced_tools():
    return list(available_tools.values())
